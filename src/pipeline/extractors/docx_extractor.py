"""
DOCX Document Extractor
=======================
Extracts text from Microsoft Word (.docx) documents using python-docx.
Extracts text from paragraphs and table cells.
"""

from pathlib import Path
from typing import List
import docx

from src.pipeline.extractors.base import BaseExtractor, ExtractionResult


class DOCXExtractor(BaseExtractor):
    """
    Extractor for Word documents (.docx) via python-docx.
    """

    def supported_extensions(self) -> List[str]:
        return [".docx"]

    def extract(self, filepath: str) -> ExtractionResult:
        path = Path(filepath)
        filename = path.name

        try:
            doc = docx.Document(filepath)
            raw_lines: List[str] = []

            # Extract from paragraphs
            for para in doc.paragraphs:
                cleaned_line = para.text.strip()
                if cleaned_line:
                    raw_lines.append(cleaned_line)

            # Extract from tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        cleaned_line = cell.text.strip()
                        if cleaned_line and cleaned_line not in raw_lines:
                            raw_lines.append(cleaned_line)

            raw_text = "\n".join(raw_lines)
            total_characters = sum(len(line) for line in raw_lines)
            total_lines = len(raw_lines)

            return ExtractionResult(
                filename=filename,
                raw_text=raw_text,
                raw_lines=raw_lines,
                total_characters=total_characters,
                total_lines=total_lines,
                success=True,
                error_message="",
            )
        except Exception as e:
            return ExtractionResult(
                filename=filename,
                raw_text="",
                raw_lines=[],
                total_characters=0,
                total_lines=0,
                success=False,
                error_message=str(e),
            )
