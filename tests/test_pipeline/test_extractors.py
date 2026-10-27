"""
Unit Tests for Document Extractors
==================================
Tests PDF, DOCX, and TXT extractors including format support, multi-encoding
fallback, page-number filtering, table cell extraction, and error handling.
"""

from pathlib import Path
import docx
import pymupdf as fitz
import pytest

from src.pipeline.extractors.docx_extractor import DOCXExtractor
from src.pipeline.extractors.pdf_extractor import PDFExtractor
from src.pipeline.extractors.txt_extractor import TXTExtractor


def test_txt_extractor_utf8(tmp_path: Path):
    """Verify TXTExtractor correctly reads UTF-8 text containing Awingi orthography."""
    file_path = tmp_path / "sample.txt"
    sample_content = "Awi agawxi qutɨrsɨ.\n\nAnbaŋ bɨsta antɨrsɨ.\n"
    file_path.write_text(sample_content, encoding="utf-8")

    extractor = TXTExtractor()
    assert extractor.supported_extensions() == [".txt"]

    res = extractor.extract(str(file_path))
    assert res.success is True
    assert res.total_lines == 2
    assert "ɨ" in res.raw_text
    assert "ŋ" in res.raw_text
    assert res.error_message == ""


def test_txt_extractor_latin1(tmp_path: Path):
    """Verify TXTExtractor fallback to latin-1 encoding."""
    file_path = tmp_path / "latin1.txt"
    sample_content = "Café au lait\nResume text\n"
    file_path.write_bytes(sample_content.encode("latin-1"))

    extractor = TXTExtractor()
    res = extractor.extract(str(file_path))
    assert res.success is True
    assert res.total_lines == 2
    assert "Café" in res.raw_text


def test_txt_extractor_nonexistent_file():
    """Verify TXTExtractor handles missing file gracefully."""
    extractor = TXTExtractor()
    res = extractor.extract("non_existent_file.txt")
    assert res.success is False
    assert res.error_message != ""
    assert res.total_lines == 0


def test_pdf_extractor_page_number_filtering(tmp_path: Path):
    """Verify PDFExtractor extracts text and discards standalone page numbers (digits <= 4)."""
    pdf_path = tmp_path / "test_doc.pdf"
    doc = fitz.open()

    # Page 1
    page1 = doc.new_page()
    page1.insert_text(
        fitz.Point(50, 72),
        "Awi agawxi qutersi.\n1\n",
    )

    # Page 2
    page2 = doc.new_page()
    page2.insert_text(
        fitz.Point(50, 72),
        "Dinxewa keda agawxi.\n25\n",
    )

    doc.save(str(pdf_path))
    doc.close()

    extractor = PDFExtractor()
    assert extractor.supported_extensions() == [".pdf"]

    res = extractor.extract(str(pdf_path))
    assert res.success is True
    # "1" and "25" should be filtered out
    assert "1" not in res.raw_lines
    assert "25" not in res.raw_lines
    assert len(res.raw_lines) == 2
    assert "Awi agawxi qutersi." in res.raw_lines
    assert "Dinxewa keda agawxi." in res.raw_lines


def test_pdf_extractor_invalid_file(tmp_path: Path):
    """Verify PDFExtractor gracefully handles corrupt/invalid PDF files."""
    corrupt_pdf = tmp_path / "corrupt.pdf"
    corrupt_pdf.write_bytes(b"not a real pdf content")

    extractor = PDFExtractor()
    res = extractor.extract(str(corrupt_pdf))
    assert res.success is False
    assert res.error_message != ""


def test_docx_extractor_paragraphs_and_tables(tmp_path: Path):
    """Verify DOCXExtractor extracts from paragraphs and tables without empty lines."""
    docx_path = tmp_path / "test_doc.docx"
    doc = docx.Document()

    # Paragraphs
    doc.add_paragraph("Awi agawxi qutɨrsɨ.")
    doc.add_paragraph("   ")  # Empty line should be skipped
    doc.add_paragraph("Anbaŋ bɨsta antɨrsɨ.")

    # Table
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Table Item One"
    table.cell(0, 1).text = "Table Item Two"
    table.cell(1, 0).text = " "  # Empty cell
    table.cell(1, 1).text = "Table Item Three"

    doc.save(str(docx_path))

    extractor = DOCXExtractor()
    assert extractor.supported_extensions() == [".docx"]

    res = extractor.extract(str(docx_path))
    assert res.success is True
    assert "Awi agawxi qutɨrsɨ." in res.raw_lines
    assert "Anbaŋ bɨsta antɨrsɨ." in res.raw_lines
    assert "Table Item One" in res.raw_lines
    assert "Table Item Two" in res.raw_lines
    assert "Table Item Three" in res.raw_lines
    assert " " not in res.raw_lines
    assert res.total_lines == 5


def test_docx_extractor_invalid_file(tmp_path: Path):
    """Verify DOCXExtractor gracefully handles corrupt/invalid docx files."""
    corrupt_docx = tmp_path / "corrupt.docx"
    corrupt_docx.write_bytes(b"not a real docx content")

    extractor = DOCXExtractor()
    res = extractor.extract(str(corrupt_docx))
    assert res.success is False
    assert res.error_message != ""
